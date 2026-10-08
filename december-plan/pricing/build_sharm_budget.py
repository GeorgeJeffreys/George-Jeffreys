"""Egypt option B: Sharm el-Sheikh (Naama Bay), Red Sea Diving College basis, first cut 3 Oct 2026 on Alain Sobol's
quote (co-owner, email 3 Oct 08:13 UTC, thread 1a0f4547d32792bd). All figures USD as quoted.

Club pricing rule (George, 23-24 Sep; beginner fee above certified fee, 30 Sep): ONE trip fee covers everything
diving-related (course incl. eLearning, boat days, all gear incl. computer, park and environmental fees) and the
hotel two sharing with breakfast, plus a share of the leader's costs (flight allowance, gear, odd-room risk) and a 6%
contingency (same basis as the Marsa Shagra cut, so the two compare like for like). Own costs: flights, transfers,
dive-accident cover, visa, lunches off the boat and dinners, drinks, tips. RSDC gives 1 free place per 10 paying.

Blue inputs on the Inputs sheet; everything else is formulas. Run, then recalc with LibreOffice.
Flags: QUOTED = RSDC email 3 Oct; ADVERTISED = public page or snippet; ASSUMED/INFERRED = our estimate.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
BLUE = Font(color="1F4E9A"); BOLD = Font(bold=True); H = Font(bold=True, size=13)
FILL = PatternFill("solid", fgColor="EEF3F8"); MONEY = '#,##0'; PCT = '0%'

ws = wb.active; ws.title = "Inputs"
for col, w in zip("ABCD", (72, 12, 12, 70)): ws.column_dimensions[col].width = w
ws['A1'] = "Sharm group budget — inputs (blue cells are editable; USD)"; ws['A1'].font = H
ws.append(["item", "value", "unit", "flag / source"])
for c in ws[2]: c.font = BOLD; c.fill = FILL
R = {}
def inp(key, label, value, unit, flag, fmt=None):
    ws.append([label, value, unit, flag]); r = ws.max_row
    ws.cell(r, 2).font = BLUE
    if fmt: ws.cell(r, 2).number_format = fmt
    R[key] = f"Inputs!$B${r}"
def h2(t):
    ws.append([t]); ws.cell(ws.max_row, 1).font = BOLD

h2("Pricing basis")
inp('n_price', "Pricing headcount: people including George at which the fee is set", 20, "people", "policy (same as Marsa Shagra cut)")
inp('n_min', "Minimum headcount the contingency must cover", 12, "people", "policy")
inp('beg_share', "Share of participants who are Open Water students", 0.5, "share", "brief", PCT)
inp('contingency', "Contingency inside the fee", 0.06, "share", "policy (Marsa Shagra basis)", PCT)
h2("Red Sea Diving College, Naama Bay (QUOTED 3 Oct 2026, Alain Sobol, per person)")
inp('room_pp', "Tao Beach Hotel, Naama Bay, two sharing, breakfast, per person per night", 70, "USD/night", "QUOTED (cheaper hotels offered on request)", MONEY)
inp('nights', "Nights, Sun 3 - Fri 8 Jan", 5, "nights", "programme")
inp('ow', "PADI Open Water: course incl. the PADI eLearning code, full gear, certification", 445, "USD", "QUOTED (eLearning code included, confirmed 3 Oct)", MONEY)
inp('boat3', "Certified boat diving, 3 days / 6 dives (Mon-Wed) incl. boat fees, lunch, drinks, environmental fee", 325, "USD", "QUOTED (Ras Mohammed and the Straits of Tiran, confirmed 3 Oct)", MONEY)
inp('thu_boat', "Thursday morning boat, 2 dives, whole group, all included", 100, "USD", "QUOTED", MONEY)
inp('thistle', "Thistlegorm supplement (AOW + 20 logged dives), taken on the third boat day", 160, "USD", "QUOTED 3 Oct; add-on, not in the fee", MONEY)
inp('hb', "Half-board supplement per person per night (dinner)", 15, "USD/night", "QUOTED 3 Oct 13:23; add-on so the beginner fee stays above certified", MONEY)
inp('aow', "Advanced Open Water, 2 days, all included", 395, "USD", "QUOTED; add-on", MONEY)
inp('rental', "Full equipment rental incl. computer, per day", 35, "USD/day", "QUOTED", MONEY)
inp('rent_days_cert', "Rental days, certified (Mon-Thu)", 4, "days", "programme")
inp('rent_days_beg', "Rental days, new divers (Thursday only; course includes gear)", 1, "days", "programme")
inp('free_every', "Free places: 1 per this many paying participants", 10, "paying", "QUOTED (an extra one 'to discuss' at size)")
h2("Leader costs carried by the group (George's hotel and diving come from the first free place)")
inp('g_flight', "George's flight allowance", 1000, "USD", "policy; PHL-SSH advertised from about 908 return (sources.csv)", MONEY)
inp('odd_room', "Odd-headcount room risk (half a place)", f"={R['room_pp']}*{R['nights']}*0.5", "USD", "ASSUMED", MONEY)
inp('g_rent', "George's gear rental, 4 days", f"={R['rental']}*4", "USD", "formula", MONEY)
h2("Own spending guidance, not in the fee (no flights)")
inp('xfer', "Airport transfers both ways, SSH to Naama Bay", 15, "USD", "QUOTED 3 Oct (approx., arranged by RSDC)", MONEY)
inp('dan_own', "Dive-accident cover, required", 84, "USD", "ADVERTISED (DAN)", MONEY)
inp('visa', "Egypt visa on arrival", 30, "USD", "DOCUMENTED (secondary)", MONEY)
inp('meals', "Dinners, and lunches on non-boat days, per day", 30, "USD/day", "INFERRED (Naama Bay restaurants)", MONEY)
inp('extras', "Drinks and nights out per day", 25, "USD/day", "INFERRED", MONEY)
inp('days', "Days in Sharm", 5, "days", "")
inp('tips', "Tips for crews and guides", 50, "USD", "INFERRED", MONEY)

fee = wb.create_sheet("Fee")
fee.column_dimensions['A'].width = 66
for col in "BC": fee.column_dimensions[col].width = 18
fee['A1'] = "Trip fee per person at the pricing headcount (USD; two sharing, breakfast)"; fee['A1'].font = H
fee.append(["", "Open Water student", "Certified diver"])
for c in fee[2]: c.font = BOLD; c.fill = FILL
paying = f"({R['n_price']}-1)"
leader = f"({R['g_flight']}+{R['odd_room']}+{R['g_rent']})"
rows = [
 ("Diving (Open Water + Thursday boat, or 3 boat days + Thursday boat)", f"={R['ow']}+{R['thu_boat']}", f"={R['boat3']}+{R['thu_boat']}"),
 ("Gear rental incl. computer", f"={R['rental']}*{R['rent_days_beg']}", f"={R['rental']}*{R['rent_days_cert']}"),
 ("Tao Beach Hotel, 5 nights, breakfast, two sharing", f"={R['room_pp']}*{R['nights']}", f"={R['room_pp']}*{R['nights']}"),
 ("Leader's costs shared (flight allowance, gear, odd-room risk)", f"={leader}/{paying}", f"={leader}/{paying}"),
 ("Subtotal", "=SUM(B3:B6)", "=SUM(C3:C6)"),
 ("Contingency", f"=ROUND(B7*{R['contingency']},0)", f"=ROUND(C7*{R['contingency']},0)"),
 ("Fee before rounding", "=B7+B8", "=C7+C8"),
 ("TRIP FEE, USD (rounded up to the nearest 50)", "=CEILING(B9,50)", "=CEILING(C9,50)"),
 ("", None, None),
 ("Optional add-ons at cost", None, None),
 ("Advanced Open Water", "n/a", f"={R['aow']}"),
 ("Thistlegorm supplement (AOW + 20 dives)", "n/a", f"={R['thistle']}"),
 ("Half board, 5 dinners", f"={R['hb']}*{R['nights']}", f"={R['hb']}*{R['nights']}"),
 ("", None, None),
 ("GUIDANCE BUDGET WITHOUT FLIGHTS", None, None),
 ("Trip fee", "=B10", "=C10"),
 ("Own spending: transfers, dive cover, visa, meals, drinks, tips", f"={R['xfer']}+{R['dan_own']}+{R['visa']}+({R['meals']}+{R['extras']})*{R['days']}+{R['tips']}", f"={R['xfer']}+{R['dan_own']}+{R['visa']}+({R['meals']}+{R['extras']})*{R['days']}+{R['tips']}"),
 ("Guidance budget", "=B18+B19", "=C18+C19"),
 ("Guidance budget, certified taking AOW and the Thistlegorm", "n/a", "=C20+C13+C14"),
]
for label, b, c in rows:
    fee.append([label, b, c]); r = fee.max_row
    for col in (2, 3): fee.cell(r, col).number_format = MONEY
for r in (7, 10, 18, 20, 21):
    for cell in fee[r]: cell.font = BOLD

g = wb.create_sheet("Group")
g.column_dimensions['A'].width = 70
cols = {"B": 12, "C": 20, "D": 25, "E": 30}
for col in cols: g.column_dimensions[col].width = 14
g['A1'] = "Group budget at four sizes, USD (half students; George leads and does not pay)"; g['A1'].font = H
g.append(["", "12 people (minimum)", "20 people (pricing basis)", "25 people", "30 people"])
for c in g[2]: c.font = BOLD; c.fill = FILL
def grow(label, fn, fmt=MONEY, bold=False):
    g.append([label] + [fn(col, n) for col, n in cols.items()]); r = g.max_row
    for i in range(2, 6): g.cell(r, i).number_format = fmt
    if bold:
        for cell in g[r]: cell.font = BOLD
    return r
r_n = grow("People including George", lambda c, n: n, '0')
r_pay = grow("Paying participants", lambda c, n: f"={c}{r_n}-1", '0')
r_beg = grow("Open Water students", lambda c, n: f"=ROUND({c}{r_pay}*{R['beg_share']},0)", '0')
r_cert = grow("Paying certified divers", lambda c, n: f"={c}{r_pay}-{c}{r_beg}", '0')
r_free = grow("Free places earned (1 per 10 paying)", lambda c, n: f"=INT({c}{r_pay}/{R['free_every']})", '0')
r_rev = grow("Trip fees collected", lambda c, n: f"={c}{r_beg}*Fee!$B$10+{c}{r_cert}*Fee!$C$10", bold=True)
r_hotel = grow("Hotel, paying guests", lambda c, n: f"={c}{r_pay}*{R['room_pp']}*{R['nights']}")
r_ow = grow("Open Water courses and Thursday boat", lambda c, n: f"={c}{r_beg}*({R['ow']}+{R['thu_boat']})")
r_bt = grow("Boat diving, paying certified", lambda c, n: f"={c}{r_cert}*({R['boat3']}+{R['thu_boat']})")
r_gpl = grow("George's hotel and diving (free if a free place is earned)", lambda c, n: f"=IF({c}{r_free}>=1,0,{R['room_pp']}*{R['nights']}+{R['boat3']}+{R['thu_boat']})")
r_spare = grow("Spare free places credited (hotel + boat package)", lambda c, n: f"=-MAX({c}{r_free}-1,0)*({R['room_pp']}*{R['nights']}+{R['boat3']}+{R['thu_boat']})")
r_rent = grow("Gear rental (certified incl. George 4 days; students 1 day)", lambda c, n: f"=({c}{r_cert}+1)*{R['rental']}*{R['rent_days_cert']}+{c}{r_beg}*{R['rental']}*{R['rent_days_beg']}")
r_odd = grow("Odd-headcount room risk", lambda c, n: f"=IF(MOD({c}{r_n},2)=1,{R['odd_room']},0)")
r_gfl = grow("George's flight allowance", lambda c, n: f"={R['g_flight']}")
r_costs = grow("TOTAL COSTS", lambda c, n: f"={c}{r_hotel}+{c}{r_ow}+{c}{r_bt}+{c}{r_gpl}+{c}{r_spare}+{c}{r_rent}+{c}{r_odd}+{c}{r_gfl}", bold=True)
r_sur = grow("Surplus (refunded pro rata or held)", lambda c, n: f"={c}{r_rev}-{c}{r_costs}", bold=True)
grow("Covered?", lambda c, n: f'=IF({c}{r_sur}>=0,"yes","NO")', '@', bold=True)
g.append(["Deposit and cancellation terms: RSDC will set them once the group size is known (QUOTED 3 Oct). Space held to the end of October (QUOTED 3 Oct). RSDC is inside the Tao hotel on the Naama Bay promenade."])

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save("sharm_group_budget.xlsx")
print("written sharm_group_budget.xlsx")
