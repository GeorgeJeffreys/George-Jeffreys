"""Cozumel group budget, Blue Note + Hotel Plaza basis (George's choice, 23 Sep 2026).

Pricing rule (George, 23 Sep evening): participants pay ONE trip fee that covers everything except their own
flight: diving, park fees, eLearning, DAN cover and the hotel with breakfast (everyone shares, George too),
plus a share of the leader's costs (George's flight allowance, his half-room and DAN), plus a
contingency. The fee is set at a conservative pricing headcount so that any larger group produces a surplus
(refunded pro rata or held as buffer); the contingency absorbs a shortfall down to the minimum headcount.
Blue Note's 10% commission on certified packages goes to George. Spare comps reduce the club's invoice.

Blue inputs on the Inputs sheet; everything else is formulas. Run, then recalc with LibreOffice.
Flags: QUOTED = Blue Note / Hotel Plaza emails 23 Sep; ADVERTISED = public page; INFERRED = our estimate.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
BLUE = Font(color="1F4E9A"); BOLD = Font(bold=True); H = Font(bold=True, size=13)
FILL = PatternFill("solid", fgColor="EEF3F8"); MONEY = '#,##0'; PCT = '0%'

# ---------------- Inputs ----------------
ws = wb.active; ws.title = "Inputs"
for col, w in zip("ABCD", (72, 12, 12, 64)): ws.column_dimensions[col].width = w
ws['A1'] = "Cozumel group budget — inputs (blue cells are editable)"; ws['A1'].font = H
ws.append(["item", "value", "unit", "flag / source"])
for c in ws[2]: c.font = BOLD; c.fill = FILL
R = {}
def inp(key, label, value, unit, flag, fmt=None):
    ws.append([label, value, unit, flag]); r = ws.max_row
    ws.cell(r, 2).font = BLUE
    if fmt: ws.cell(r, 2).number_format = fmt
    R[key] = f"Inputs!$B${r}"
def h2(t):
    ws.append([t]); ws.cell(ws.max_row, 1).font = BOLD

h2("Pricing basis")
inp('n_price', "Pricing headcount: people including George at which the fee is set (larger groups make a surplus)", 20, "people", "policy: conservative")
inp('n_min', "Minimum headcount the contingency must cover", 12, "people", "policy")
inp('beg_share', "Share of participants who are Open Water students", 0.5, "share", "brief", PCT)
inp('contingency', "Contingency inside the fee (refunded pro rata if unused)", 0.06, "share", "policy", PCT)
h2("Blue Note Scuba (QUOTED 23 Sep 2026, Doug Atkinson)")
inp('ow', "Open Water course, 3 days, equipment included", 400, "USD", "QUOTED", MONEY)
inp('ow_extra', "Two extra boat days for new divers (Thu–Fri)", 240, "USD", "QUOTED", MONEY)
inp('pkg', "Certified 5-day two-tank package", 600, "USD", "QUOTED", MONEY)
inp('rental', "Equipment rental per day, everyone rents (in the fee)", 30, "USD/day", "QUOTED", MONEY)
inp('rent_days_cert', "Rental days, certified", 5, "days", "programme")
inp('rent_days_beg', "Rental days, new divers (course includes equipment; the two extra boat days do not)", 2, "days", "ASSUMED; ask Blue Note whether the USD 240 extra days include gear")
inp('park', "Marine park fee per boat day", 15, "USD/day", "ADVERTISED; whether it is inside the package is asked", MONEY)
inp('park_days_cert', "Boat days, certified", 5, "days", "programme")
inp('park_days_beg', "Boat days, new divers (dive 3–4 day + Thu + Fri)", 3, "days", "programme")
inp('comm', "Commission on certified packages, paid to George", 0.10, "share", "QUOTED: deducted from the invoice", PCT)
inp('comp_per', "One free certified package per N paid", 5, "paid", "QUOTED")
inp('dep_bn', "Blue Note deposit at booking", 0.25, "share", "QUOTED; balance 30 days out; refund to 7 days out", PCT)
inp('pm2', "Afternoon two-tank dive (optional add-on)", 130, "USD", "ADVERTISED price page", MONEY)
inp('night_charter', "Private night charter, up to 10 divers (optional add-on, shared)", 800, "USD", "ADVERTISED price page", MONEY)
inp('nitrox_fill', "Nitrox fill per tank (course free with eLearning)", 12, "USD", "INFERRED; QUOTE NEEDED", MONEY)
h2("Hotel Plaza Cozumel (QUOTED via Blue Note's contracted rate)")
inp('room_lo', "Room per night, 2 people with breakfast, low-rate nights", 70, "USD", "QUOTED (Dec 14–15 rate)", MONEY)
inp('room_hi', "Room per night, 2 people with breakfast, high-rate nights", 88, "USD", "QUOTED (Dec 16–19 rate)", MONEY)
inp('nights_lo', "Low-rate nights (Sun 13, Mon 14)", 2, "nights", "assumes the Dec 14–15 rate covers the nights of 13 and 14; CONFIRM")
inp('nights_hi', "High-rate nights (Tue 15 – Fri 18)", 4, "nights", "")
inp('hotel_dep', "Hotel deposit at booking (share of block)", 0.25, "share", "QUOTE NEEDED; assumption", PCT)
h2("Included in the fee for everyone")
inp('elearn_ow', "PADI Open Water eLearning (club buys the codes)", 195, "USD", "ADVERTISED", MONEY)
inp('dan', "DAN membership + Preferred plan, bought for each person", 119, "USD", "ADVERTISED", MONEY)
inp('tax_buffer', "Tax buffer per head (Mexican IVA and lodging tax if the quotes are pre-tax)", 200, "USD", "policy (George, 23 Sep): assume USD 200 until Doug confirms", MONEY)
h2("Leader costs carried by the group (George shares a room like everyone)")
inp('g_flight', "George's flight allowance", 700, "USD", "policy; PHL–CZM band 568–800", MONEY)
inp('g_room', "George's half of a shared room, 6 nights", f"=({R['room_lo']}*{R['nights_lo']}+{R['room_hi']}*{R['nights_hi']})/2", "USD", "formula", MONEY)
inp('odd_room', "Odd-headcount room risk (one person alone at the group's cost)", f"={R['g_room']}", "USD", "formula: half a room", MONEY)
inp('g_dan', "George's DAN cover", 119, "USD", "ADVERTISED", MONEY)
inp('g_rent', "George's gear rental, 5 days", f"={R['rental']}*5", "USD", "formula", MONEY)
h2("Own spending guidance, not in the fee (no flights)")
inp('meals', "Lunches, dinners and drinks per day (breakfast included)", 45, "USD/day", "INFERRED for an MBA crowd", MONEY)
inp('days_meals', "Days on the island", 6, "days", "")
inp('tips', "Tips for boat crew and hotel", 50, "USD", "INFERRED", MONEY)
inp('xfer', "Airport transfers both ways", 30, "USD", "ADVERTISED shared van 15 each way", MONEY)
inp('elearn_aow', "PADI AOW eLearning (add-on for AOW takers)", 220, "USD", "ADVERTISED (about)", MONEY)

# ---------------- Fee ----------------
fee = wb.create_sheet("Fee")
fee.column_dimensions['A'].width = 62
for col in "BC": fee.column_dimensions[col].width = 18
fee['A1'] = "Trip fee per person, set at the pricing headcount (USD, 2 sharing)"; fee['A1'].font = H
fee.append(["", "Open Water student", "Certified diver"])
for c in fee[2]: c.font = BOLD; c.fill = FILL
room_pp = f"(({R['room_lo']}*{R['nights_lo']}+{R['room_hi']}*{R['nights_hi']})/2)"
paying = f"({R['n_price']}-1)"
leader_total = f"({R['g_flight']}+{R['g_room']}+{R['odd_room']}+{R['g_dan']}+{R['g_rent']})"
rows = [
 ("Diving with Blue Note", f"={R['ow']}+{R['ow_extra']}", f"={R['pkg']}"),
 ("Gear rental, everyone", f"={R['rental']}*{R['rent_days_beg']}", f"={R['rental']}*{R['rent_days_cert']}"),
 ("Marine park fees", f"={R['park']}*{R['park_days_beg']}", f"={R['park']}*{R['park_days_cert']}"),
 ("Tax buffer (IVA and lodging tax, if the quotes are pre-tax)", f"={R['tax_buffer']}", f"={R['tax_buffer']}"),
 ("PADI eLearning", f"={R['elearn_ow']}", "0"),
 ("DAN dive-accident cover", f"={R['dan']}", f"={R['dan']}"),
 ("Hotel Plaza, 6 nights with breakfast, 2 sharing", f"={room_pp}", f"={room_pp}"),
 ("Leader's costs shared (George's flight, half-room, DAN, rental, odd-room risk)", f"={leader_total}/{paying}", f"={leader_total}/{paying}"),
 ("Subtotal", "=SUM(B3:B10)", "=SUM(C3:C10)"),
 ("Contingency", f"=ROUND(B11*{R['contingency']},0)", f"=ROUND(C11*{R['contingency']},0)"),
 ("TRIP FEE (rounded up to the nearest 100, as posted)", "=CEILING(B11+B12,100)", "=CEILING(C11+C12,100)"),
 ("", None, None),
 ("Optional add-ons at sign-up, charged at cost", None, None),
 ("AOW: course free with eLearning; eLearning", "n/a", f"={R['elearn_aow']}"),
 ("Nitrox: course free with eLearning; 10 fills", "n/a", f"={R['nitrox_fill']}*10"),
 ("Afternoon two-tank dive, each, incl. park fee", f"={R['pm2']}+{R['park']}", f"={R['pm2']}+{R['park']}"),
 ("Night dive, private charter shared by 10", f"={R['night_charter']}/10", f"={R['night_charter']}/10"),
 ("", None, None),
 ("GUIDANCE BUDGET WITHOUT FLIGHTS", None, None),
 ("Trip fee", "=B13", "=C13"),
 ("Own spending: meals, drinks, tips, airport transfers", f"={R['meals']}*{R['days_meals']}+{R['tips']}+{R['xfer']}", f"={R['meals']}*{R['days_meals']}+{R['tips']}+{R['xfer']}"),
 ("Guidance budget", "=B22+B23", "=C22+C23"),
 ("Guidance budget, certified taking AOW and nitrox", "n/a", "=C24+C16+C17"),
]
for label, b, c in rows:
    fee.append([label, b, c]); r = fee.max_row
    for col in (2, 3): fee.cell(r, col).number_format = MONEY
for r in (11, 13, 21, 24, 25):
    for cell in fee[r]: cell.font = BOLD
fee['A27'] = "The fee covers everything diving-related (course or package, gear, park fees, eLearning), DAN insurance and the hotel; everything else is the person's own. It is set at the pricing headcount; a bigger group makes a surplus that is refunded pro rata or held as buffer, a smaller group down to the minimum is covered by the contingency (see Group sheet). George's place: diving covered by Blue Note's comp, flight and room carried by the group, commission to George, all disclosed on the sign-up page."

# ---------------- Group ----------------
g = wb.create_sheet("Group")
g.column_dimensions['A'].width = 66
cols = {"B": 12, "C": 20, "D": 30, "E": 40}
for col in cols: g.column_dimensions[col].width = 14
g['A1'] = "Group budget at four sizes (half students; George leads and does not pay)"; g['A1'].font = H
g.append(["", "12 people (minimum)", "20 people (pricing basis)", "30 people", "40 people"])
for c in g[2]: c.font = BOLD; c.fill = FILL
def grow(label, fn, fmt=MONEY, bold=False):
    g.append([label] + [fn(col, n) for col, n in cols.items()]); r = g.max_row
    for i in range(2, 6): g.cell(r, i).number_format = fmt
    if bold:
        for cell in g[r]: cell.font = BOLD
    return r
def sect(t):
    g.append([t]); g.cell(g.max_row, 1).font = BOLD
r_n = grow("People including George", lambda c, n: n, '0')
r_pay = grow("Paying participants", lambda c, n: f"={c}{r_n}-1", '0')
r_beg = grow("Open Water students", lambda c, n: f"=ROUND({c}{r_pay}*{R['beg_share']},0)", '0')
r_cert = grow("Paying certified divers", lambda c, n: f"={c}{r_pay}-{c}{r_beg}", '0')
r_rooms = grow("Rooms, everyone 2 sharing including George (odd person alone at the group's cost)", lambda c, n: f"=ROUNDUP({c}{r_n}/2,0)", '0')
sect("Revenue")
r_rev = grow("Trip fees collected", lambda c, n: f"={c}{r_beg}*Fee!$B$13+{c}{r_cert}*Fee!$C$13", bold=True)
sect("Costs")
r_bn_ow = grow("Blue Note: Open Water courses and extra days", lambda c, n: f"={c}{r_beg}*({R['ow']}+{R['ow_extra']})")
r_bn_pk = grow("Blue Note: certified packages, paid divers", lambda c, n: f"={c}{r_cert}*{R['pkg']}")
r_comps = grow("Comps earned (1 per N paid certified)", lambda c, n: f"=INT({c}{r_cert}/{R['comp_per']})", '0')
r_gdive = grow("George's package (free if at least one comp)", lambda c, n: f"=IF({c}{r_comps}>=1,0,{R['pkg']})")
r_spare = grow("Spare comps credited against the invoice", lambda c, n: f"=-MAX({c}{r_comps}-1,0)*{R['pkg']}")
r_bn_gross = grow("Blue Note gross invoice", lambda c, n: f"={c}{r_bn_ow}+{c}{r_bn_pk}+{c}{r_gdive}+{c}{r_spare}")
r_comm = grow("Less 10% commission on paid certified packages", lambda c, n: f"=-{c}{r_bn_pk}*{R['comm']}")
r_bn_net = grow("Blue Note net payable", lambda c, n: f"={c}{r_bn_gross}+{c}{r_comm}", bold=True)
r_comm_out = grow("Commission passed to George", lambda c, n: f"=-{c}{r_comm}")
r_rent = grow("Gear rental for everyone including George", lambda c, n: f"={c}{r_beg}*{R['rental']}*{R['rent_days_beg']}+({c}{r_cert}+1)*{R['rental']}*{R['rent_days_cert']}")
r_tax = grow("Tax buffer held (paid to Blue Note and the hotel if charged; otherwise refunded)", lambda c, n: f"={c}{r_n}*{R['tax_buffer']}")
r_park = grow("Marine park fees (cash on the island, paid from the float)", lambda c, n: f"={c}{r_beg}*{R['park']}*{R['park_days_beg']}+({c}{r_cert}+1)*{R['park']}*{R['park_days_cert']}")
r_el = grow("PADI Open Water eLearning codes", lambda c, n: f"={c}{r_beg}*{R['elearn_ow']}")
r_dan = grow("DAN cover for everyone including George", lambda c, n: f"={c}{r_n}*{R['dan']}")
r_hotel = grow("Hotel Plaza block, all rooms", lambda c, n: f"={c}{r_rooms}*({R['room_lo']}*{R['nights_lo']}+{R['room_hi']}*{R['nights_hi']})")
r_gfl = grow("George's flight allowance", lambda c, n: f"={R['g_flight']}")
r_costs = grow("TOTAL COSTS", lambda c, n: f"={c}{r_bn_net}+{c}{r_comm_out}+{c}{r_rent}+{c}{r_tax}+{c}{r_park}+{c}{r_el}+{c}{r_dan}+{c}{r_hotel}+{c}{r_gfl}", bold=True)
sect("Result")
r_sur = grow("Surplus (contingency and over-recovery; refunded pro rata or held)", lambda c, n: f"={c}{r_rev}-{c}{r_costs}", bold=True)
r_surpp = grow("Surplus per paying person", lambda c, n: f"={c}{r_sur}/{c}{r_pay}")
r_surpct = grow("Surplus as a share of revenue", lambda c, n: f"={c}{r_sur}/{c}{r_rev}", PCT)
r_cover = grow("Covered? (surplus must stay positive at the minimum headcount)", lambda c, n: f'=IF({c}{r_sur}>=0,"yes","NO")', '@', bold=True)
r_pot = grow("George's pot: commission (flight and room already carried above)", lambda c, n: f"={c}{r_comm_out}", bold=True)
sect("Deposits and cash")
r_dep_in = grow("Participant deposits at sign-up (USD 400 each)", lambda c, n: f"={c}{r_pay}*400")
r_dep_out = grow("Blue Note 25% deposit + hotel deposit", lambda c, n: f"={c}{r_bn_gross}*{R['dep_bn']}+{c}{r_hotel}*{R['hotel_dep']}")
r_gap = grow("Deposit cover (positive means deposits fund the outflow)", lambda c, n: f"={c}{r_dep_in}-{c}{r_dep_out}", bold=True)
r_bal_in = grow("Balances due from participants by Fri 6 Nov", lambda c, n: f"={c}{r_rev}-{c}{r_dep_in}")
r_bal_out = grow("Blue Note balance Fri 13 Nov + rental + hotel balance + eLearning + DAN", lambda c, n: f"={c}{r_bn_net}-{c}{r_bn_gross}*{R['dep_bn']}+{c}{r_rent}+{c}{r_hotel}*(1-{R['hotel_dep']})+{c}{r_el}+{c}{r_dan}")
r_float = grow("Cash float to carry to the island (park fees)", lambda c, n: f"={c}{r_park}")
g.append(["Notes: revenue is the fixed fee from the Fee sheet, so the surplus grows with headcount (the leader share is over-recovered) and shrinks below the pricing basis; the contingency keeps it positive down to the minimum. Spare comps reduce the invoice for the group; the commission goes to George."])

# ---------------- Payments ----------------
pay = wb.create_sheet("Payments")
pay.column_dimensions['A'].width = 22; pay.column_dimensions['B'].width = 76; pay.column_dimensions['C'].width = 24
pay['A1'] = "Payment schedule and refund policy (mirrors Blue Note's terms)"; pay['A1'].font = H
for row in [
 ("date", "what", "who"),
 ("by Fri 2 Oct", "Sign-ups open on CampusGroups; deposit USD 400 per person holds a place", "participants"),
 ("Sat 10 Oct", "Sign-up deadline for the first block; headcount to Blue Note and Hotel Plaza", "George"),
 ("by Thu 15 Oct", "Blue Note deposit 25% of the gross invoice; hotel deposit (terms to confirm)", "club account"),
 ("Fri 6 Nov", "Participant balances due (fee less deposit, plus any add-ons)", "participants"),
 ("Fri 13 Nov", "Blue Note final payment (30 days out); hotel balance per its terms; buy eLearning codes and DAN", "club account"),
 ("Sun 6 Dec", "Blue Note's 7-day refund cut-off", "everyone"),
 ("Sun 13 Dec", "Arrival; park fees paid from the club float; meals, drinks, tips and transfers are each person's own", "George"),
 ("after the trip", "Unused contingency and any surplus refunded pro rata", "club account"),
 ("", "", ""),
 ("Refund ladder", "Deposit refundable in full until Fri 6 Nov; 7 Nov to 6 Dec refundable only if the place is resold or Blue Note refunds it; from 7 Dec no refund, only Blue Note credit where given. Weather cancellations: Blue Note reschedules or refunds the missed dives.", "policy"),
 ("Money handling", "All payments through CampusGroups or Blue Note's own link, never personal Venmo or Zelle. The sign-up page states that the leader's diving is comped by the shop, his flight and half-room are carried by the group, and the shop's commission goes to him.", "policy"),
]:
    pay.append(list(row))
for c in pay[2]: c.font = BOLD; c.fill = FILL

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
out = "/home/user/George-Jeffreys/december-plan/pricing/cozumel_group_budget.xlsx"
wb.save(out); print("saved", out)
