"""Egypt group budget, Red Sea Diving Safari (Marsa Shagra) basis, first cut 24 Sep 2026, second cut 26 Sep on Riad
Tarek's reservation answers (email 07:50 UTC, thread 1a0ca71a4aee743e): five-night rate, four one-day passes instead of
the five-day package, rental with computer, 20% non-refundable deposit, cancellation ladder, low season, meet-and-assist.

Club pricing rule (George, 23–24 Sep): ONE trip fee covers everything diving-related (course or package, all gear
rental, park and environmental fees, e-code) and the lodging with full board (everyone shares, George too), plus a
share of the leader's costs (George's flight allowance, gear, odd-room risk) and a 6% contingency. Own costs:
flights, airport transfers, dive-accident cover (DAN or equivalent, required), visa, tips, drinks. No tax buffer
(RSDS prices are all-in in EUR). RSDS gives the 11th, 21st and 30th paying guest free (accommodation + diving +
transfer): the first free place is George's; spare free places reduce the invoice. No commission is offered.

Blue inputs on the Inputs sheet; everything else is formulas. Run, then recalc with LibreOffice.
Flags: QUOTED = RSDS emails 24 and 26 Sep; ADVERTISED = public page; ASSUMED = our estimate.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
BLUE = Font(color="1F4E9A"); BOLD = Font(bold=True); H = Font(bold=True, size=13)
FILL = PatternFill("solid", fgColor="EEF3F8"); MONEY = '#,##0'; PCT = '0%'; FX = '0.0000'

ws = wb.active; ws.title = "Inputs"
for col, w in zip("ABCD", (72, 12, 12, 70)): ws.column_dimensions[col].width = w
ws['A1'] = "Egypt group budget — inputs (blue cells are editable; EUR unless stated)"; ws['A1'].font = H
ws.append(["item", "value", "unit", "flag / source"])
for c in ws[2]: c.font = BOLD; c.fill = FILL
R = {}
def inp(key, label, value, unit, flag, fmt=None):
    ws.append([label, value, unit, flag]); r = ws.max_row
    ws.cell(r, 2).font = BLUE
    if fmt: ws.cell(r, 2).number_format = fmt
    R[key] = f"Inputs!$B${r}"
def h2(t):
    ws.append([t]); ws.cell(ws.max_row, 1).font = BOLD

h2("Pricing basis")
inp('fx', "USD per EUR", 1.1465, "USD/EUR", "ADVERTISED mid-market 21 Sep 2026 (sources.csv)", FX)
inp('n_price', "Pricing headcount: people including George at which the fee is set", 20, "people", "policy: conservative")
inp('n_min', "Minimum headcount the contingency must cover", 12, "people", "policy (11 paying = first free place)")
inp('beg_share', "Share of participants who are Open Water students", 0.5, "share", "brief", PCT)
inp('contingency', "Contingency inside the fee (refunded pro rata if unused)", 0.06, "share", "policy", PCT)
h2("Red Sea Diving Safari, Marsa Shagra (QUOTED 24 Sep 2026, Riad Tarek, low season, per person)")
inp('room_pp', "Royal Tent, double, full board, per person per night", 78, "EUR/night", "QUOTED 26 Sep (390 for five nights; Deluxe Chalet 113/night = 565)", MONEY)
inp('nights', "Nights, Sun 3 – Fri 8 Jan", 5, "nights", "programme")
inp('pkg', "House reef diving, four one-day unlimited passes Mon–Thu (certified), EUR 77 a day", 308, "EUR", "QUOTED 26 Sep (no four-day package; cheaper than the five-day 347)", MONEY)
inp('day_extra', "One-day unlimited house-reef pass (the new divers' Thursday after certification)", 77, "EUR", "QUOTED 26 Sep", MONEY)
inp('ow', "Open Water course incl. equipment", 358, "EUR", "QUOTED", MONEY)
inp('ecode', "PADI e-code / certification per beginner", 94, "EUR", "QUOTED", MONEY)
inp('aow', "Advanced Open Water incl. equipment (add-on)", 286, "EUR", "QUOTED (242 with own gear)", MONEY)
inp('excursion', "Dolphin House day boat (certified, Wed) incl. park fee", 67, "EUR", "QUOTED 48 + 19", MONEY)
inp('elph', "Elphinstone speedboat dive (add-on, AOW)", 36, "EUR/dive", "QUOTED", MONEY)
inp('env', "Environmental fee 10 + hyperbaric fee 7, once per person", 17, "EUR", "QUOTED", MONEY)
inp('rental', "Full equipment rental incl. computer per day, certified divers (course prices include gear)", 39, "EUR/day", "QUOTED 26 Sep (33 without computer; club rule: computers inside the fee)", MONEY)
inp('rent_days_cert', "Rental days, certified (Mon–Thu)", 4, "days", "programme")
inp('rent_days_beg', "Rental days, new divers (Thursday only; course includes gear)", 1, "days", "programme")
inp('free_every', "Free place thresholds: 11th, 21st and 30th paying guest (accommodation + diving + transfer)", 11, "paying", "QUOTED; modelled as 11, 21, 30")
inp('dep_rsds', "RSDS deposit at booking (non-refundable)", 0.20, "share", "QUOTED 26 Sep; balance on site, 5% off for cash, card +3%", PCT)
h2("Leader costs carried by the group (George's lodging and diving come from the first free place)")
inp('g_flight', "George's flight allowance", 1000, "USD", "policy; PHL–HRG band 700–1,300 (sources.csv)", MONEY)
inp('odd_room', "Odd-headcount room risk (one person alone at the group's cost, single supplement unknown)", f"={R['room_pp']}*{R['nights']}*0.5", "EUR", "ASSUMED half a place", MONEY)
inp('g_rent', "George's gear rental, 4 days", f"={R['rental']}*4", "EUR", "formula", MONEY)
h2("Own spending guidance, not in the fee (no flights)")
inp('xfer', "Airport transfers both ways, Hurghada, 10–20 pax rate", 82, "EUR", "QUOTED 41 each way (33 at 21–30 pax; Marsa Alam 22 / 17)", MONEY)
inp('dan_own', "Dive-accident cover, required, bought by each person", 84, "USD", "ADVERTISED: DAN membership 40 + Master 44 (short-term plan about 50 if sold locally)", MONEY)
inp('visa', "Egypt visa on arrival", 30, "USD", "DOCUMENTED (secondary): USD 30 since 1 Mar 2026", MONEY)
inp('assist', "Visa on arrival plus meet-and-assist at the airport (optional, minimum four)", 45, "EUR", "QUOTED 26 Sep", MONEY)
inp('extras', "Drinks, snacks, tips per day (full board covers meals)", 20, "USD/day", "INFERRED", MONEY)
inp('days_extras', "Days at the village", 5, "days", "")
inp('tips', "Tips for guides and staff", 50, "USD", "INFERRED", MONEY)

# ---------------- Fee ----------------
fee = wb.create_sheet("Fee")
fee.column_dimensions['A'].width = 66
for col in "BC": fee.column_dimensions[col].width = 18
fee['A1'] = "Trip fee per person, set at the pricing headcount (EUR lines, USD fee; 2 sharing, full board)"; fee['A1'].font = H
fee.append(["", "Open Water student", "Certified diver"])
for c in fee[2]: c.font = BOLD; c.fill = FILL
paying = f"({R['n_price']}-1)"
leader_eur = f"({R['g_flight']}/{R['fx']}+{R['odd_room']}+{R['g_rent']})"
rows = [
 ("Diving with RSDS (course + e-code + Thursday pass, or four day passes Mon–Thu)", f"={R['ow']}+{R['ecode']}+{R['day_extra']}", f"={R['pkg']}"),
 ("Dolphin House day boat incl. park fee (certified, Wed)", "0", f"={R['excursion']}"),
 ("Gear rental", f"={R['rental']}*{R['rent_days_beg']}", f"={R['rental']}*{R['rent_days_cert']}"),
 ("Environmental and hyperbaric fees", f"={R['env']}", f"={R['env']}"),
 ("Royal Tent, 5 nights full board, 2 sharing", f"={R['room_pp']}*{R['nights']}", f"={R['room_pp']}*{R['nights']}"),
 ("Leader's costs shared (George's flight allowance, gear, odd-room risk)", f"={leader_eur}/{paying}", f"={leader_eur}/{paying}"),
 ("Subtotal, EUR", "=SUM(B3:B8)", "=SUM(C3:C8)"),
 ("Contingency, EUR", f"=ROUND(B9*{R['contingency']},0)", f"=ROUND(C9*{R['contingency']},0)"),
 ("Fee in USD before rounding", f"=(B9+B10)*{R['fx']}", f"=(C9+C10)*{R['fx']}"),
 ("TRIP FEE, USD (rounded up to the nearest 50, as posted)", "=CEILING(B11,50)", "=CEILING(C11,50)"),
 ("", None, None),
 ("Optional add-ons at sign-up, charged at cost (USD)", None, None),
 ("Advanced Open Water incl. equipment", "n/a", f"={R['aow']}*{R['fx']}"),
 ("Elphinstone speedboat dive, each", "n/a", f"={R['elph']}*{R['fx']}"),
 ("Abu Dabbab afternoon (48)", f"=48*{R['fx']}", f"=48*{R['fx']}"),
 ("", None, None),
 ("GUIDANCE BUDGET WITHOUT FLIGHTS (USD)", None, None),
 ("Trip fee", "=B12", "=C12"),
 ("Own spending: transfers, dive cover, visa, meet-and-assist, drinks, tips", f"=({R['xfer']}+{R['assist']})*{R['fx']}+{R['dan_own']}+{R['visa']}+{R['extras']}*{R['days_extras']}+{R['tips']}", f"=({R['xfer']}+{R['assist']})*{R['fx']}+{R['dan_own']}+{R['visa']}+{R['extras']}*{R['days_extras']}+{R['tips']}"),
 ("Guidance budget", "=B20+B21", "=C20+C21"),
 ("Guidance budget, certified taking AOW and one Elphinstone dive", "n/a", "=C22+C15+C16"),
]
for label, b, c in rows:
    fee.append([label, b, c]); r = fee.max_row
    for col in (2, 3): fee.cell(r, col).number_format = MONEY
for r in (9, 12, 20, 22, 23):
    for cell in fee[r]: cell.font = BOLD
fee['A25'] = "The fee covers everything diving-related (course, e-code, package, gear, fees) and five nights full board at Marsa Shagra; flights, transfers, dive-accident cover, visa, drinks and tips are each person's own. Set at the pricing headcount; a bigger group makes a surplus that is refunded pro rata or held; the contingency covers a shortfall down to the minimum. George's place: lodging and diving from RSDS's first free place, flight allowance carried by the group, all disclosed on the sign-up page."

# ---------------- Group ----------------
g = wb.create_sheet("Group")
g.column_dimensions['A'].width = 70
cols = {"B": 12, "C": 20, "D": 25, "E": 30}
for col in cols: g.column_dimensions[col].width = 14
g['A1'] = "Group budget at four sizes, EUR unless stated (half students; George leads and does not pay)"; g['A1'].font = H
g.append(["", "12 people (minimum)", "20 people (pricing basis)", "25 people", "30 people"])
for c in g[2]: c.font = BOLD; c.fill = FILL
def grow(label, fn, fmt=MONEY, bold=False):
    g.append([label] + [fn(col, n) for col, n in cols.items()]); r = g.max_row
    for i in range(2, 6): g.cell(r, i).number_format = fmt
    if bold:
        for cell in g[r]: cell.font = BOLD
    return r
def sect(t):
    g.append([t]); g.cell(g.max_row, 1).font = BOLD
r_n = grow("People including George", lambda c, n: n, '0')
r_pay = grow("Paying participants", lambda c, n: f"={c}{r_n}-1", '0')
r_beg = grow("Open Water students", lambda c, n: f"=ROUND({c}{r_pay}*{R['beg_share']},0)", '0')
r_cert = grow("Paying certified divers", lambda c, n: f"={c}{r_pay}-{c}{r_beg}", '0')
r_free = grow("Free places earned (11th, 21st, 30th paying guest)", lambda c, n: f"=({c}{r_pay}>=11)+({c}{r_pay}>=21)+({c}{r_pay}>=30)", '0')
sect("Revenue (USD fees converted to EUR)")
r_rev = grow("Trip fees collected", lambda c, n: f"=({c}{r_beg}*Fee!$B$12+{c}{r_cert}*Fee!$C$12)/{R['fx']}", bold=True)
sect("Costs")
r_lodge = grow("Lodging, paying guests, full board", lambda c, n: f"={c}{r_pay}*{R['room_pp']}*{R['nights']}")
r_ow = grow("Open Water courses, e-codes and Thursday day", lambda c, n: f"={c}{r_beg}*({R['ow']}+{R['ecode']}+{R['day_extra']})")
r_pk = grow("House reef packages and Dolphin House, paying certified", lambda c, n: f"={c}{r_cert}*({R['pkg']}+{R['excursion']})")
r_gplace = grow("George's lodging and package (free if a free place is earned)", lambda c, n: f"=IF({c}{r_free}>=1,0,{R['room_pp']}*{R['nights']}+{R['pkg']})")
r_spare = grow("Spare free places credited against the invoice (lodging + package)", lambda c, n: f"=-MAX({c}{r_free}-1,0)*({R['room_pp']}*{R['nights']}+{R['pkg']})")
r_env = grow("Environmental and hyperbaric fees, everyone", lambda c, n: f"={c}{r_n}*{R['env']}")
r_rsds = grow("RSDS invoice", lambda c, n: f"={c}{r_lodge}+{c}{r_ow}+{c}{r_pk}+{c}{r_gplace}+{c}{r_spare}+{c}{r_env}", bold=True)
r_rent = grow("Gear rental (certified incl. George 4 days; students 1 day)", lambda c, n: f"=({c}{r_cert}+1)*{R['rental']}*{R['rent_days_cert']}+{c}{r_beg}*{R['rental']}*{R['rent_days_beg']}")
r_odd = grow("Odd-headcount room risk", lambda c, n: f"=IF(MOD({c}{r_n},2)=1,{R['odd_room']},0)")
r_gfl = grow("George's flight allowance (EUR)", lambda c, n: f"={R['g_flight']}/{R['fx']}")
r_costs = grow("TOTAL COSTS", lambda c, n: f"={c}{r_rsds}+{c}{r_rent}+{c}{r_odd}+{c}{r_gfl}", bold=True)
sect("Result")
r_sur = grow("Surplus, EUR (contingency and over-recovery; refunded pro rata or held)", lambda c, n: f"={c}{r_rev}-{c}{r_costs}", bold=True)
r_surusd = grow("Surplus, USD", lambda c, n: f"={c}{r_sur}*{R['fx']}")
r_cover = grow("Covered? (surplus must stay positive at the minimum headcount)", lambda c, n: f'=IF({c}{r_sur}>=0,"yes","NO")', '@', bold=True)
r_pot = grow("George's pot, EUR: value of spare free places (no commission offered)", lambda c, n: f"=-{c}{r_spare}", bold=True)
sect("Deposits and cash (RSDS: 20% non-refundable at booking, balance on site)")
r_dep_in = grow("Participant deposits at sign-up (USD 400 each, in EUR)", lambda c, n: f"={c}{r_pay}*400/{R['fx']}")
r_dep_out = grow("RSDS deposit (share of invoice)", lambda c, n: f"={c}{r_rsds}*{R['dep_rsds']}")
r_gap = grow("Deposit cover (positive means deposits fund the outflow)", lambda c, n: f"={c}{r_dep_in}-{c}{r_dep_out}", bold=True)
g.append(["Notes: fees are posted in USD and converted at the workbook rate, so a weaker dollar eats the surplus first; the contingency and the 12-person floor are the cushion. A free place beyond George's own is credited to the club here; George can instead ask to keep it as his pot (a TDI course, parked). Paying the balance in cash on site earns a 5% discount on the RSDS invoice, not modelled."])

# ---------------- Payments ----------------
pay = wb.create_sheet("Payments")
pay.column_dimensions['A'].width = 22; pay.column_dimensions['B'].width = 80; pay.column_dimensions['C'].width = 24
pay['A1'] = "Payment schedule (RSDS terms quoted 26 Sep; dates provisional)"; pay['A1'].font = H
for row in [
 ("date", "what", "who"),
 ("by Fri 9 Oct", "Sign-ups open on CampusGroups; deposit USD 400 per person holds a place", "participants"),
 ("Fri 16 Oct", "Sign-up deadline; headcount and course split to RSDS", "George"),
 ("by Fri 23 Oct", "RSDS deposit, 20% of the invoice, non-refundable, by bank transfer or card; e-codes issued to students", "club account"),
 ("Fri 13 Nov", "Participant balances due (fee less deposit, plus add-ons)", "participants"),
 ("Fri 8 Jan, on site", "RSDS balance paid on departure: cash in EUR earns 5% off, card in EGP costs 3% (George carries the club card or a bank draft; decide in December)", "George / club account"),
 ("Sun 3 Jan", "Arrival; transfers, visa and dive cover are each person's own", "George"),
 ("after the trip", "Unused contingency and any surplus refunded pro rata", "club account"),
 ("", "", ""),
 ("Refund ladder", "Mirrors RSDS: deposit non-refundable once paid to RSDS (23 Oct); cancellation more than 30 days out costs 20% of the place, 15–29 days 30%, 7–14 days 50%, 6 days or less 80%, no-show 100%; a resold place is refunded in full less the club's costs.", "policy"),
 ("Cancellation ladder (RSDS)", ">30 days 20% / 15–29 days 30% / 7–14 days 50% / <=6 days 80% / no-show 100%; low season 3–8 Jan, no supplements", "QUOTED 26 Sep"),
 ("Money handling", "All payments through CampusGroups or RSDS's own link, never personal Venmo or Zelle. The sign-up page states that the leader's place comes from the village's free-place rule and his flight allowance is carried by the group.", "policy"),
]:
    pay.append(list(row))
for c in pay[2]: c.font = BOLD; c.fill = FILL

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save("egypt_group_budget.xlsx")
print("written egypt_group_budget.xlsx")
